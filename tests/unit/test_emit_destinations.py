"""Emit preserves percent-encoding; classify still uses the de-obfuscated view."""

from __future__ import annotations

import io
import re
import zipfile
import zlib
from html import escape

import pytest

from artifactsmith.renderers import get_renderer
from artifactsmith.renderers.html_sanitize import sanitize_html_document
from artifactsmith.renderers.links import classify_href, public_href_or_none
from artifactsmith.renderers.md_sanitize import collect_link_destinations, sanitize_markdown
from artifactsmith.renderers.safety import sanitize_text

# (name, input_url, expected_emit)
_PRESERVE_CASES: list[tuple[str, str, str]] = [
    ("slash_in_path", "https://example.com/a%2Fb", "https://example.com/a%2Fb"),
    ("hash_in_path", "https://example.com/wiki/Foo%23Bar", "https://example.com/wiki/Foo%23Bar"),
    ("space_in_path", "https://example.com/q%20w", "https://example.com/q%20w"),
    (
        "query_space_and_amp",
        "https://example.com/search?q=a%20b&x=1",
        "https://example.com/search?q=a%20b&x=1",
    ),
    (
        "parens_wikipedia",
        "https://en.wikipedia.org/wiki/Foo_(bar)",
        "https://en.wikipedia.org/wiki/Foo_(bar)",
    ),
    ("fragment", "https://example.com/path#frag", "https://example.com/path#frag"),
    (
        "encoded_slash_query_fragment",
        "https://example.com/a%2Fb?x=1&y=2#sec",
        "https://example.com/a%2Fb?x=1&y=2#sec",
    ),
    ("idn_unicode", "https://bücher.example/x", "https://bücher.example/x"),
    (
        "protocol_relative_keeps_encoding",
        "//www.nsf.gov/a%2Fb",
        "https://www.nsf.gov/a%2Fb",
    ),
]


@pytest.mark.parametrize(
    "name,raw,expected",
    _PRESERVE_CASES,
    ids=[c[0] for c in _PRESERVE_CASES],
)
def test_public_href_or_none_preserves_encoding(name, raw, expected):
    assert public_href_or_none(raw) == expected


def test_literal_space_emit_becomes_percent20():
    assert public_href_or_none("https://example.com/has space") == "https://example.com/has%20space"


@pytest.mark.parametrize(
    "name,raw,expected",
    _PRESERVE_CASES,
    ids=[c[0] for c in _PRESERVE_CASES],
)
def test_markdown_emit_exact_destination(name, raw, expected):
    out = sanitize_markdown(f"[label]({raw})")
    dests = collect_link_destinations(out)
    if name == "idn_unicode":
        # markdown-it punycodes IDN before tokens reach us.
        assert dests == ["https://xn--bcher-kva.example/x"], (out, dests)
        expected = dests[0]
    else:
        assert dests == [expected], (name, out, dests)
    if "(" in expected or ")" in expected:
        assert f"](<{expected}>)" in out
    else:
        assert f"]({expected})" in out


@pytest.mark.parametrize(
    "name,raw,expected",
    [c for c in _PRESERVE_CASES if not c[1].startswith("//")],
    ids=[c[0] for c in _PRESERVE_CASES if not c[1].startswith("//")],
)
def test_html_emit_exact_destination(name, raw, expected):
    # Put & as &amp; in HTML source so the attribute value keeps a real &.
    href_attr = escape(raw, quote=True)
    doc = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'><title>t</title></head>"
        f"<body><p><a href={href_attr}>label</a></p></body></html>"
    )
    out = sanitize_html_document(doc)
    # HTML attribute serialization escapes & as &amp; (required format escaping).
    attr_expected = escape(expected, quote=True)
    assert f'href="{attr_expected}"' in out


def _pdf_uris(pdf: bytes) -> list[str]:
    """Extract /URI destinations; support PDF hex strings and literal parentheses escapes."""
    uris: list[str] = []
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", pdf, re.S):
        try:
            dec = zlib.decompress(m.group(1))
        except zlib.error:
            continue
        for u in re.findall(rb"/URI\s*\((?:\\.|[^\\)])*\)", dec):
            inner = u[u.find(b"(") + 1 : -1]
            # Unescape PDF string escapes \( \) \\
            text = (
                inner.replace(b"\\(", b"(")
                .replace(b"\\)", b")")
                .replace(b"\\\\", b"\\")
                .decode("utf-8", errors="replace")
            )
            uris.append(text)
        for u in re.findall(rb"/URI\s*<([0-9A-Fa-f\s]+)>", dec):
            hexes = re.sub(rb"\s+", b"", u)
            try:
                uris.append(bytes.fromhex(hexes.decode("ascii")).decode("utf-8", errors="replace"))
            except ValueError:
                continue
    return uris


@pytest.mark.parametrize(
    "name,raw,expected",
    [c for c in _PRESERVE_CASES if not c[1].startswith("//")],
    ids=[c[0] for c in _PRESERVE_CASES if not c[1].startswith("//")],
)
def test_pdf_docx_xlsx_emit_exact_destination(name, raw, expected):
    body = sanitize_text(f"See [label]({raw}) end.")
    if name == "idn_unicode":
        expected_opts = {"https://bücher.example/x", "https://xn--bcher-kva.example/x"}
    else:
        expected_opts = {expected}

    pdf = get_renderer("pdf").render(title="T", body=body).files["document.pdf"]
    pdf_uris = set(_pdf_uris(pdf))
    assert pdf_uris & expected_opts, (name, pdf_uris)

    docx = get_renderer("docx").render(title="T", body=body).files["document.docx"]
    with zipfile.ZipFile(io.BytesIO(docx)) as zf:
        rels = "\n".join(zf.read(n).decode("utf-8", errors="replace") for n in zf.namelist() if n.endswith(".rels"))
    # OOXML may entity-escape & as &amp; in Target attributes.
    assert any(
        e in rels or escape(e, quote=True)[1:-1] in rels or e.replace("&", "&amp;") in rels for e in expected_opts
    ), (
        name,
        rels,
    )

    from openpyxl import load_workbook

    xlsx = get_renderer("xlsx").render(title="T", body=body).files["document.xlsx"]
    wb = load_workbook(io.BytesIO(xlsx))
    targets: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink is not None:
                    targets.append(getattr(cell.hyperlink, "target", None) or str(cell.hyperlink))
    assert any(t in expected_opts for t in targets), (name, targets)


def test_classify_still_blocks_percent_encoded_dangerous_scheme():
    assert classify_href("java%E2%80%8Bscript:alert(1)") == "blocked"
    assert public_href_or_none("java%E2%80%8Bscript:alert(1)") is None
    assert public_href_or_none("/\\evil.example") is None


def test_xlsx_system_prompt_steers_one_url_per_cell():
    from artifactsmith import builder as b

    prompt = b.system_prompt_for("xlsx")
    assert "own table cell" in prompt
    assert "one link per cell" in prompt
    assert "own table cell" not in b.system_prompt_for("pdf")
    assert b.system_prompt_for("pdf") == b.CONTENT_SYSTEM
