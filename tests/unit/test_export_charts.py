"""PDF, DOCX, and XLSX draw chart and diagram fences instead of printing them."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from openpyxl import load_workbook
from PIL import Image

from artifactsmith import builder
from artifactsmith.renderers import get_renderer
from artifactsmith.renderers.figures import rasterize_svg
from artifactsmith.renderers.md_parse import blocks_to_simple_html, parse_blocks

CHART = """```chart bar unit=$
June: 10
July: 20
August: 15
September: 30
```
"""

SEQUENCE = """```sequence
Client -> Server: SYN
Server --> Client: SYN-ACK
```
"""

MERMAID_CHART = """```mermaid
xychart-beta
title "Revenue"
x-axis [June, July, August, September]
bar [5, 6, 7, 8]
```
"""

MERMAID_FLOW = """```mermaid
flowchart TD
    A[June] --> B[July]
    B -->|yes| C{August}
```
"""

MERMAID_SEQUENCE = """```mermaid
sequenceDiagram
    Alice->>Bob: Hello there
```
"""

BAD_CHART = """```chart bar
Later: soon
```
"""

PYTHON = """```python
print("keep this")
```
"""


def _html(body: str) -> str:
    return blocks_to_simple_html("Months", parse_blocks(body))


def _xlsx_text(data: bytes) -> str:
    book = load_workbook(io.BytesIO(data))
    parts: list[str] = []
    for name in book.sheetnames:
        parts.append(name)
        for row in book[name].iter_rows(values_only=True):
            parts.extend("" if value is None else str(value) for value in row)
    return "\n".join(parts)


def _docx_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return package.read("word/document.xml").decode("utf-8")


def test_content_prompt_teaches_our_fences_and_refuses_mermaid():
    text = builder.CONTENT_SYSTEM
    assert "```chart bar|line|pie|donut```" in text
    assert "```sequence```" in text
    assert "```flow```" in text
    assert "```timeline```" in text
    assert "Oct 10: Blog topics approved" in text
    assert "*date | title | note" in text
    assert "Mermaid is not supported" in text
    assert builder.system_prompt_for("pdf") == text
    assert builder.system_prompt_for("docx") == text


def test_chart_and_diagram_become_figures_not_code():
    chart = parse_blocks(CHART)
    assert chart[0].kind == "figure"
    assert chart[0].figure == "chart"
    assert "<svg" in chart[0].svg
    assert "June: 10" not in chart[0].svg
    html = _html(CHART)
    assert "<figure" in html
    assert "June" in html
    assert "<pre>" not in html
    assert "June: 10" not in html

    diagram = parse_blocks(SEQUENCE)
    assert diagram[0].kind == "figure"
    assert diagram[0].figure == "diagram"
    assert "Client -&gt; Server" not in diagram[0].svg
    assert "->" not in _html(SEQUENCE)


def test_undrawable_fence_is_a_table_and_python_stays_code():
    blocks = parse_blocks(BAD_CHART)
    assert blocks[0].kind == "table"
    html = _html(BAD_CHART)
    assert "<table>" in html
    assert "<pre>" not in html
    assert "Later: soon" not in html
    assert "Later" in html and "soon" in html

    code = parse_blocks(PYTHON)
    assert code[0].kind == "code"
    assert "<pre>" in _html(PYTHON)
    assert "keep this" in _html(PYTHON)


def test_mermaid_xychart_and_flowchart_are_drawn_and_sequence_is_a_table():
    chart_html = _html(MERMAID_CHART)
    assert "<svg" in chart_html
    assert "June" in chart_html
    assert "xychart" not in chart_html
    assert "<pre>" not in chart_html

    flow_html = _html(MERMAID_FLOW)
    assert "<svg" in flow_html
    assert "June" in flow_html
    assert "July" in flow_html
    assert "flowchart" not in flow_html
    assert "-->" not in flow_html
    assert "<pre>" not in flow_html

    sequence_html = _html(MERMAID_SEQUENCE)
    assert "<table>" in sequence_html
    assert "<svg" not in sequence_html
    assert "<pre>" not in sequence_html
    assert "sequenceDiagram" not in sequence_html
    assert "->>" not in sequence_html
    assert "Alice" in sequence_html
    assert "Hello there" in sequence_html


TIMELINE = """```timeline
Oct 10: Blog topics approved
Oct 20: Publish 2 holiday ordering guides
```
"""

UNDRAWN = """```timeline
Blog topics approved
Publish the holiday guide
```
"""


def test_colon_timeline_and_undrawable_lines_survive_html_pdf_and_docx():
    """A date: text timeline draws, and a timeline that cannot draw keeps its lines."""
    from artifactsmith.layout.render import render_layout

    page, _notes = render_layout(TIMELINE, style="house", title="Next steps")
    assert "<svg" in page
    assert "Oct 10" in page and "Blog topics approved" in page
    assert "Oct 20" in page and "Publish 2 holiday ordering guides" in page
    assert "Diagram could not be drawn" not in page

    html = _html(TIMELINE)
    assert "<svg" in html
    assert "Blog topics approved" in html and "Publish 2 holiday ordering guides" in html
    assert "Diagram could not be drawn" not in html
    sizes = _replaced_sizes(html)
    assert sizes and sizes[0][0] > 40 and sizes[0][1] > 40
    pdf = get_renderer("pdf").render(title="Next steps", body=TIMELINE).files["document.pdf"]
    assert pdf.startswith(b"%PDF")
    assert b"Oct 10: Blog topics approved" not in pdf

    docx = get_renderer("docx").render(title="Next steps", body=TIMELINE).files["document.docx"]
    xml = _docx_xml(docx)
    assert "Oct 10" in xml and "Blog topics approved" in xml
    assert "Publish 2 holiday ordering guides" in xml
    assert "Diagram could not be drawn" not in xml

    listed, notes = render_layout(UNDRAWN, style="house", title="Next steps")
    assert "timeline diagram shown as a list" in notes
    assert "Blog topics approved" in listed and "Publish the holiday guide" in listed
    assert "Diagram could not be drawn" not in listed

    plain = _html(UNDRAWN)
    assert "<table>" in plain
    assert "Blog topics approved" in plain and "Publish the holiday guide" in plain
    assert "Diagram could not be drawn" not in plain
    undrawn_pdf = get_renderer("pdf").render(title="Next steps", body=UNDRAWN).files["document.pdf"]
    assert undrawn_pdf.startswith(b"%PDF")
    undrawn_xml = _docx_xml(get_renderer("docx").render(title="Next steps", body=UNDRAWN).files["document.docx"])
    assert "Blog topics approved" in undrawn_xml and "Publish the holiday guide" in undrawn_xml
    assert "Diagram could not be drawn" not in undrawn_xml


def test_script_in_a_chart_label_is_escaped():
    body = "```chart bar\nJune: 1\n<script>: 2\n```\n"
    html = _html(body)
    assert "<script" not in html
    assert "&lt;script&gt;" in html


def _replaced_sizes(html: str) -> list[tuple[float, float]]:
    from weasyprint import HTML

    from artifactsmith.renderers.pdf import DenyAllURLFetcher

    document = HTML(string=html, base_url=".", url_fetcher=DenyAllURLFetcher()).render()
    found: list[tuple[float, float]] = []

    def walk(box: object) -> None:
        if type(box).__name__ == "InlineReplacedBox":
            found.append((float(getattr(box, "width", 0)), float(getattr(box, "height", 0))))
        for child in getattr(box, "children", None) or ():
            walk(child)

    for page in document.pages:  # type: ignore[attr-defined]
        walk(page._page_box)
    return found


def test_pdf_draws_the_chart_and_docx_embeds_a_png_with_alt_text():
    html = _html(CHART)
    sizes = _replaced_sizes(html)
    assert sizes and sizes[0][0] > 100 and sizes[0][1] > 100
    pdf = get_renderer("pdf").render(title="Months", body=CHART).files["document.pdf"]
    assert pdf.startswith(b"%PDF")
    assert b"June: 10" not in pdf
    assert b"xychart" not in pdf

    data = get_renderer("docx").render(title="Months", body=CHART).files["document.docx"]
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        media = [name for name in package.namelist() if name.startswith("word/media/")]
        assert len(media) == 1
        assert package.read(media[0]).startswith(b"\x89PNG")
    xml = _docx_xml(data)
    assert "descr=" in xml
    assert "June" in xml
    assert "June: 10" not in xml
    assert "```" not in xml

    doc = Document(io.BytesIO(data))
    assert doc.inline_shapes


def test_docx_diagram_png_and_spreadsheet_table_omit_the_fence():
    data = get_renderer("docx").render(title="Handshake", body=SEQUENCE).files["document.docx"]
    xml = _docx_xml(data)
    assert "descr=" in xml
    assert "Client" in xml
    assert "->" not in xml
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        media = [name for name in package.namelist() if name.startswith("word/media/")]
        assert package.read(media[0]).startswith(b"\x89PNG")

    sheet = _xlsx_text(get_renderer("xlsx").render(title="Months", body=CHART).files["document.xlsx"])
    assert "June" in sheet
    assert "June: 10" not in sheet
    assert "```" not in sheet

    flow = _xlsx_text(get_renderer("xlsx").render(title="Flow", body=MERMAID_FLOW).files["document.xlsx"])
    assert "June" in flow
    assert "flowchart" not in flow
    assert "-->" not in flow

    bad = _xlsx_text(get_renderer("xlsx").render(title="Bad", body=BAD_CHART).files["document.xlsx"])
    assert "Later" in bad and "soon" in bad
    assert "```" not in bad


def test_rasterized_chart_is_a_colored_png_and_garbage_svg_falls_back():
    svg = parse_blocks(CHART)[0].svg
    png = rasterize_svg(svg)
    assert png is not None
    image = Image.open(io.BytesIO(png)).convert("RGB")
    colors = {color for _count, color in image.getcolors(image.size[0] * image.size[1]) or []}
    assert (11, 92, 171) in colors
    assert rasterize_svg("<svg></not>") is None
