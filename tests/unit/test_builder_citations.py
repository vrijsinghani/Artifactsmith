"""Builder path: system prompts permit public citations; fake model → linked exports."""

from __future__ import annotations

import io
import re
import zipfile
import zlib

import pytest

from artifactsmith import builder as b

# Typical model answer with labelled public links and an outside image.
_MODEL_MD = """# Sources brief

Cite the [National Science Foundation](https://www.nsf.gov/) for the program overview.

![NSF logo](https://www.nsf.gov/images/logo.png)

Pilot store code is HARBOR-17.
"""

_MODEL_HTML = """<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Sources</title>
<style>body{font-family:system-ui,sans-serif}</style></head><body>
<h1>Sources brief</h1>
<p>Cite the <a href="https://www.nsf.gov/">National Science Foundation</a>.</p>
<img src="https://www.nsf.gov/images/logo.png" alt="NSF logo">
<p>Pilot store code is HARBOR-17.</p>
</body></html>
"""

_README_CREATE = {
    "display_name": "Pilot store brief",
    "verbatim_request": "One page that states the pilot store code and why it matters.",
    "source_content": (
        "The pilot store code is HARBOR-17. It matters because it is the single identifier "
        "used across inventory, support, and rollout reports for the pilot."
    ),
    "format": "html",
}


def _pdf_uris(pdf: bytes) -> list[bytes]:
    uris: list[bytes] = []
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", pdf, re.S):
        try:
            dec = zlib.decompress(m.group(1))
        except zlib.error:
            continue
        uris.extend(re.findall(rb"/URI\s*\(([^)]+)\)", dec))
    return uris


@pytest.mark.parametrize("fmt", ["markdown", "pdf", "docx", "xlsx"])
def test_content_system_prompt_permits_public_links(fmt):
    system = b.system_prompt_for(fmt)
    assert system is b.CONTENT_SYSTEM or system == b.CONTENT_SYSTEM
    low = system.lower()
    assert "labelled markdown links" in low or "labeled markdown links" in low
    assert "public http(s)" in low
    assert "never invent urls" in low
    assert "no http(s) or ftp urls" not in low
    assert "no images that fetch remotely" not in low


def test_html_system_prompt_permits_public_citation_links():
    system = b.system_prompt_for("html")
    low = system.lower()
    assert "public http(s)" in low
    assert "<a href=" in system or "a href=" in low
    assert "any http(s) network reference" not in low


async def _fake_md_call(model, system, user, timeout=300):
    assert "labelled Markdown links" in system or "public http(s)" in system.lower()
    return (
        "===ASSUMPTIONS===\n- none\n===SUMMARY===\nSources brief with citations.\n"
        f"===FILE: content.md===\n{_MODEL_MD}\n===END===\n"
    )


async def _fake_html_call(model, system, user, timeout=300):
    assert "public http(s)" in system.lower()
    return (
        "===ASSUMPTIONS===\n- none\n===SUMMARY===\nSources brief.\n"
        f"===FILE: index.html===\n{_MODEL_HTML}\n===END===\n"
    )


@pytest.mark.asyncio
async def test_run_build_markdown_keeps_labelled_links_and_image_link(monkeypatch):
    monkeypatch.setattr(b.llm, "call", _fake_md_call)
    res = await b.run_build(
        kind="web_static",
        slug="sources",
        display_name="Sources",
        verbatim="Cite NSF and show the logo; state HARBOR-17.",
        model="test-model",
        base_source=None,
        base_version=None,
        history=[],
        progress=lambda _m: None,
        source={"source_content": "Pilot store code is HARBOR-17. NSF site: https://www.nsf.gov/"},
        format="markdown",
    )
    text = res.files["document.md"].decode()
    assert "[National Science Foundation](https://www.nsf.gov/)" in text
    assert "![NSF logo]" not in text
    assert "[NSF logo](https://www.nsf.gov/images/logo.png)" in text


@pytest.mark.asyncio
@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx"])
async def test_run_build_formats_keep_clickable_public_links(fmt, monkeypatch):
    monkeypatch.setattr(b.llm, "call", _fake_md_call)
    res = await b.run_build(
        kind="web_static",
        slug="sources",
        display_name="Sources",
        verbatim="Cite NSF and show the logo; state HARBOR-17.",
        model="test-model",
        base_source=None,
        base_version=None,
        history=[],
        progress=lambda _m: None,
        source={"source_content": "Pilot store code is HARBOR-17."},
        format=fmt,
    )
    assert b"HARBOR-17" in res.files["content.md"]
    body_md = res.files["content.md"].decode()
    assert "[National Science Foundation](https://www.nsf.gov/)" in body_md
    assert "[NSF logo](https://www.nsf.gov/images/logo.png)" in body_md

    if fmt == "pdf":
        uris = _pdf_uris(res.files["document.pdf"])
        assert b"https://www.nsf.gov/" in uris
        assert b"https://www.nsf.gov/images/logo.png" in uris
    elif fmt == "docx":
        with zipfile.ZipFile(io.BytesIO(res.files["document.docx"])) as zf:
            rels = "\n".join(
                zf.read(n).decode("utf-8", errors="replace") for n in zf.namelist() if n.endswith(".rels")
            )
        assert "https://www.nsf.gov/" in rels
        assert "https://www.nsf.gov/images/logo.png" in rels
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(res.files["document.xlsx"]))
        targets: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.hyperlink is not None:
                        targets.append(getattr(cell.hyperlink, "target", None) or str(cell.hyperlink))
        assert any("nsf.gov" in t for t in targets), targets


@pytest.mark.asyncio
async def test_run_build_html_keeps_public_link_and_rewrites_img(monkeypatch):
    monkeypatch.setattr(b.llm, "call", _fake_html_call)
    res = await b.run_build(
        kind="web_static",
        slug="sources",
        display_name="Sources",
        verbatim="Cite NSF.",
        model="test-model",
        base_source=None,
        base_version=None,
        history=[],
        progress=lambda _m: None,
        source={"source_content": "Pilot store code is HARBOR-17."},
        format="html",
    )
    html = res.files["index.html"].decode()
    assert 'href="https://www.nsf.gov/"' in html
    assert "<img" not in html.lower()
    assert 'href="https://www.nsf.gov/images/logo.png"' in html


@pytest.mark.asyncio
async def test_readme_create_example_builds_with_supporting_source(monkeypatch):
    """Exact README create payload: source text must support 'why it matters'."""

    async def fake_call(model, system, user, timeout=300):
        assert "HARBOR-17" in user
        assert "why it matters" in user.lower() or "matters" in user.lower()
        # Model can answer because source explains why.
        assert "single identifier" in user or "HARBOR-17" in user
        return (
            "===ASSUMPTIONS===\n- none\n===SUMMARY===\nPilot brief.\n"
            "===FILE: index.html===\n<!DOCTYPE html><html><head><meta charset='utf-8'>"
            "<title>Pilot</title></head><body><h1>Pilot store brief</h1>"
            "<p>The pilot store code is HARBOR-17. It matters because it is the single "
            "identifier used across inventory, support, and rollout reports.</p>"
            "</body></html>\n===END===\n"
        )

    monkeypatch.setattr(b.llm, "call", fake_call)
    res = await b.run_build(
        kind="web_static",
        slug="pilot-store-brief",
        display_name=_README_CREATE["display_name"],
        verbatim=_README_CREATE["verbatim_request"],
        model="test-model",
        base_source=None,
        base_version=None,
        history=[],
        progress=lambda _m: None,
        source={"source_content": _README_CREATE["source_content"]},
        format="html",
    )
    html = res.files["index.html"].decode()
    assert "HARBOR-17" in html
    assert "matters" in html.lower()
