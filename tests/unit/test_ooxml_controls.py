"""C0 controls (except tab/LF/CR) must not appear in any OOXML package part."""

from __future__ import annotations

import io
import re
import zipfile

import pytest

from artifactsmith.renderers import get_renderer
from artifactsmith.renderers.safety import strip_ooxml_controls

_FORBIDDEN_C0_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("hello\x0bworld", "helloworld"),  # VT
        ("hello\x0cworld", "helloworld"),  # FF
        ("a\x00b\x01c", "abc"),
        ("keep\t\n\rthese", "keep\t\n\rthese"),
    ],
)
def test_strip_ooxml_controls(raw, expected):
    assert strip_ooxml_controls(raw) == expected


def _assert_zip_xml_clean(data: bytes) -> None:
    assert data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        xml_parts = [n for n in zf.namelist() if n.endswith(".xml") or n.endswith(".rels")]
        assert xml_parts, "expected XML parts in package"
        for name in xml_parts:
            text = zf.read(name).decode("utf-8", errors="replace")
            bad = _FORBIDDEN_C0_RE.findall(text)
            assert not bad, f"{name} contains forbidden C0: {bad!r}"


@pytest.mark.parametrize("ctrl", ["\x0b", "\x0c", "\x00", "\x01", "\x1f"])
def test_docx_strips_c0_in_title_and_body_all_xml_parts(ctrl: str):
    title = f"Ti{ctrl}tle"
    body = f"Line with{ctrl}control and keep\ttab."
    out = get_renderer("docx").render(title=title, body=body)
    _assert_zip_xml_clean(out.files["document.docx"])
    with zipfile.ZipFile(io.BytesIO(out.files["document.docx"])) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    assert "Line with" in xml
    assert "control" in xml
    assert ctrl not in xml


@pytest.mark.parametrize("ctrl", ["\x0b", "\x0c", "\x00", "\x01", "\x1f"])
def test_xlsx_strips_c0_in_title_and_body_all_xml_parts(ctrl: str):
    title = f"Sh{ctrl}eet"
    body = f"| h |\n| --- |\n| cell{ctrl}value |\n\nNote {ctrl}line."
    out = get_renderer("xlsx").render(title=title, body=body)
    _assert_zip_xml_clean(out.files["document.xlsx"])
    # Sheet name path uses stripped title; empty-after-strip still falls back.
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    assert wb.sheetnames
    joined = " ".join(
        str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert ctrl not in joined
    assert "cell" in joined or "Note" in joined


def test_xlsx_title_only_controls_still_builds_with_fallback_name():
    out = get_renderer("xlsx").render(title="\x0b\x0c", body="just text")
    _assert_zip_xml_clean(out.files["document.xlsx"])
