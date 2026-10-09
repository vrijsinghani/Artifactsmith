"""Deterministic renderers: each format produces expected bytes; public links kept."""

from __future__ import annotations

import io
import zipfile

from artifactsmith.renderers import SUPPORTED_FORMATS, get_renderer

MD = """# Pilot store

The pilot store code is HARBOR-17.

## Numbers

| store | code |
| --- | --- |
| pilot | HARBOR-17 |
| next | HARBOR-24 |

- keep the fact
"""

MD_WITH_LINK = """# Sources

See [the paper](https://example.com/paper) and https://example.org/notes.
"""


def test_supported_formats():
    assert SUPPORTED_FORMATS == frozenset({"html", "markdown", "pdf", "docx", "xlsx"})


def test_html_passthrough():
    body = "<html><body><p>HARBOR-17</p></body></html>"
    out = get_renderer("html").render(title="t", body=body)
    assert out.primary == "index.html"
    assert b"HARBOR-17" in out.files["index.html"]


def test_markdown_titles_when_needed():
    out = get_renderer("markdown").render(title="Pilot", body="The code is HARBOR-17.")
    text = out.files["document.md"].decode()
    assert text.startswith("# Pilot")
    assert "HARBOR-17" in text
    assert out.primary == "document.md"


def test_pdf_magic_and_source():
    out = get_renderer("pdf").render(title="Pilot", body=MD)
    assert out.files["document.pdf"].startswith(b"%PDF")
    assert b"HARBOR-17" in out.files["content.md"]
    assert out.primary == "document.pdf"


def test_pdf_embeds_public_link_annotation_without_fetch():
    """WeasyPrint writes /URI annotations; DenyAllURLFetcher means nothing is fetched."""
    import re
    import zlib

    out = get_renderer("pdf").render(title="Sources", body=MD_WITH_LINK)
    pdf = out.files["document.pdf"]
    assert pdf.startswith(b"%PDF")
    uris: list[bytes] = []
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", pdf, re.S):
        try:
            dec = zlib.decompress(m.group(1))
        except zlib.error:
            continue
        uris.extend(re.findall(rb"/URI\s*\(([^)]+)\)", dec))
    assert b"https://example.com/paper" in uris
    assert b"https://example.org/notes" in uris


def test_pdf_url_fetcher_denies_all():
    from artifactsmith.renderers.pdf import DenyAllURLFetcher

    try:
        DenyAllURLFetcher().fetch("https://example.com/x.png")
        raise AssertionError("expected deny-all fetcher to raise")
    except ValueError as e:
        assert "denied" in str(e).lower()


def test_weasyprint_meets_object_fetcher_floor():
    import weasyprint

    major = int(str(weasyprint.__version__).split(".", 1)[0])
    assert major >= 70, f"WeasyPrint {weasyprint.__version__} is below the object URLFetcher floor (70)"


def test_docx_is_ooxml_zip():
    out = get_renderer("docx").render(title="Pilot", body=MD)
    data = out.files["document.docx"]
    assert data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
    assert any(n.startswith("word/") for n in names)
    assert b"HARBOR-17" in out.files["content.md"]


def test_docx_has_external_hyperlinks():
    out = get_renderer("docx").render(title="Sources", body=MD_WITH_LINK)
    data = out.files["document.docx"]
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        rels = "\n".join(zf.read(n).decode("utf-8", errors="replace") for n in zf.namelist() if n.endswith(".rels"))
    assert "https://example.com/paper" in rels
    assert "https://example.org/notes" in rels
    assert 'TargetMode="External"' in rels


def test_xlsx_has_workbook():
    out = get_renderer("xlsx").render(title="Pilot", body=MD)
    data = out.files["document.xlsx"]
    assert data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        assert "[Content_Types].xml" in zf.namelist()
    assert out.primary == "document.xlsx"


def test_xlsx_public_urls_are_plain_hyperlinks_not_formulas():
    from openpyxl import load_workbook

    out = get_renderer("xlsx").render(title="Sources", body=MD_WITH_LINK)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    found = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                if "example.com" in text or "example.org" in text:
                    found = True
                    assert cell.data_type == "s"
                    assert not text.startswith("=")
                    assert cell.hyperlink is not None
                    target = getattr(cell.hyperlink, "target", None) or str(cell.hyperlink)
                    assert target.startswith("http")
    assert found


def test_markdown_and_pdf_rewrite_remote_image_to_link():
    from artifactsmith.renderers.safety import sanitize_text

    body = sanitize_text("See ![chart](https://cdn.example.com/plot.png) in the brief.")
    assert "![chart]" not in body
    assert "[chart](https://cdn.example.com/plot.png)" in body
    md = get_renderer("markdown").render(title="T", body=body)
    assert b"[chart](https://cdn.example.com/plot.png)" in md.files["document.md"]
    pdf = get_renderer("pdf").render(title="T", body=body)
    assert pdf.files["document.pdf"].startswith(b"%PDF")


def test_xlsx_multi_url_cell_keeps_all_text_links_first():
    """First public URL is the cell hyperlink; every URL stays visible as text."""
    from openpyxl import load_workbook

    body = "See https://example.com/a and https://example.org/b together."
    out = get_renderer("xlsx").render(title="Multi", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    found = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                if "example.com/a" in text and "example.org/b" in text:
                    found = True
                    assert cell.hyperlink is not None
                    target = getattr(cell.hyperlink, "target", None) or str(cell.hyperlink)
                    assert "example.com/a" in target
                    assert "example.org" not in target
    assert found


def test_xlsx_without_table_still_writes_sheet():
    out = get_renderer("xlsx").render(title="Notes only", body="Just a sentence with HARBOR-17.")
    assert out.files["document.xlsx"][:2] == b"PK"


def test_xlsx_formula_like_cells_are_literal_strings():
    """Untrusted '=' cells must remain strings after save-and-reload, not formulas."""
    import io

    from openpyxl import load_workbook

    md = """| name | value |
| --- | --- |
| a | =1+2 |
| b | =HYPERLINK("http://evil","x") |

=CMD('calc')
"""
    out = get_renderer("xlsx").render(title="=1+1", body=md)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    found_eq = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                if text.startswith("="):
                    found_eq = True
                    assert cell.data_type == "s", f"{ws.title}!{cell.coordinate} is {cell.data_type}"
                    assert getattr(cell, "data_type", None) == "s"
    assert found_eq, "expected at least one leading-= cell in the fixture"
