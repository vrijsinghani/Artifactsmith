"""PDF/DOCX/XLSX render bugs: emphasis, code wrap, numeric cells, sheet names."""

from __future__ import annotations

import io
import re
import zipfile

import pytest

from artifactsmith.renderers import get_renderer
from artifactsmith.renderers.md_parse import blocks_to_simple_html, parse_blocks
from artifactsmith.renderers.pdf import DenyAllURLFetcher

LONG_CMD = (
    "docker compose exec -T server artifactsmith token add "
    "abcdefghijklmnopqrstuvwxyz0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ "
    "--workspace alpha"
)
UNBROKEN = "X" * 240


def _iter_boxes(box: object) -> list[object]:
    out = [box]
    for child in getattr(box, "children", None) or ():
        out.extend(_iter_boxes(child))
    return out


def _layout(html: str) -> object:
    from weasyprint import HTML

    return HTML(string=html, base_url=".", url_fetcher=DenyAllURLFetcher()).render()


def _text_boxes(document: object) -> list[object]:
    boxes: list[object] = []
    for page in document.pages:  # type: ignore[attr-defined]
        page_box = getattr(page, "_page_box", None)
        if page_box is None:
            continue
        boxes.extend(b for b in _iter_boxes(page_box) if getattr(b, "text", None))
    return boxes


def test_pdf_and_docx_render_inline_emphasis():
    body = "Use **bold**, *italic*, and `inline code` in the brief."
    html = blocks_to_simple_html("T", parse_blocks(body))
    assert "<strong>bold</strong>" in html
    assert "<em>italic</em>" in html
    assert "<code>inline code</code>" in html
    assert "**bold**" not in html

    document = _layout(html)
    texts = [str(getattr(box, "text", "")) for box in _text_boxes(document)]
    joined = "".join(texts)
    assert "bold" in joined
    assert "italic" in joined
    assert "inline code" in joined
    assert "**" not in joined
    by_text = {str(getattr(box, "text", "")): box for box in _text_boxes(document)}
    assert int(by_text["bold"].style["font_weight"]) >= 700
    assert str(by_text["italic"].style["font_style"]) == "italic"
    assert "monospace" in " ".join(str(f) for f in by_text["inline code"].style["font_family"])

    pdf = get_renderer("pdf").render(title="T", body=body).files["document.pdf"]
    assert pdf.startswith(b"%PDF")
    assert b"**bold**" not in pdf

    from docx import Document

    data = get_renderer("docx").render(title="T", body=body).files["document.docx"]
    doc = Document(io.BytesIO(data))
    runs = [run for p in doc.paragraphs for run in p.runs]
    joined_docx = "".join(run.text for run in runs)
    assert "bold" in joined_docx
    assert "italic" in joined_docx
    assert "inline code" in joined_docx
    assert "**" not in joined_docx
    assert any(run.bold and "bold" in run.text for run in runs)
    assert any(run.italic and "italic" in run.text for run in runs)
    assert any(run.font.name == "Courier New" and "inline code" in run.text for run in runs)

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "<w:b" in xml
    assert "<w:i" in xml


def test_pdf_wraps_long_code_line():
    body = f"```\n{LONG_CMD}\n{UNBROKEN}\n```"
    html = blocks_to_simple_html("T", parse_blocks(body))
    assert "white-space:pre-wrap" in html
    assert "overflow-wrap:anywhere" in html
    assert "word-break:break-all" in html

    document = _layout(html)
    texts = [str(getattr(box, "text", "")) for box in _text_boxes(document)]
    joined = "".join(texts)
    assert "docker compose exec" in joined
    assert "--workspace alpha" in joined
    assert UNBROKEN in joined.replace("\n", "")

    code_lines = [t for t in texts if "docker" in t or "workspace" in t or set(t) == {"X"}]
    assert len(code_lines) >= 2

    for page in document.pages:  # type: ignore[attr-defined]
        page_box = getattr(page, "_page_box", None)
        if page_box is None:
            continue
        page_width = float(page.width)
        for box in _iter_boxes(page_box):
            x = float(getattr(box, "position_x", 0) or 0)
            w = float(getattr(box, "width", 0) or 0)
            assert x + w <= page_width + 1.5, (x, w, page_width)

    pdf = get_renderer("pdf").render(title="T", body=body).files["document.pdf"]
    assert pdf.startswith(b"%PDF")


def test_xlsx_numeric_cell_types_and_formats():
    from openpyxl import load_workbook

    body = """## Metrics

| label | value |
| --- | --- |
| count | 1234 |
| grouped | 1,234 |
| decimal | 12.50 |
| money | $1,234.56 |
| pounds | £100 |
| pct | 12% |
| pct_dec | 12.5% |
| id | 007 |
| code | HARBOR-17 |
| formula | =1+2 |
"""
    out = get_renderer("xlsx").render(title="Metrics", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    ws = next(s for s in wb.worksheets if s.title != "Notes")
    by_label = {str(row[0].value): row[1] for row in ws.iter_rows(min_row=2) if row[0].value is not None}

    count = by_label["count"]
    assert count.value == 1234
    assert count.data_type == "n"

    grouped = by_label["grouped"]
    assert grouped.value == 1234
    assert grouped.data_type == "n"
    assert "," in grouped.number_format

    decimal = by_label["decimal"]
    assert float(decimal.value) == 12.5
    assert decimal.data_type == "n"
    assert "0.00" in decimal.number_format

    money = by_label["money"]
    assert float(money.value) == 1234.56
    assert money.data_type == "n"
    assert "$" in money.number_format

    pounds = by_label["pounds"]
    assert pounds.value == 100
    assert pounds.data_type == "n"
    assert "£" in pounds.number_format

    pct = by_label["pct"]
    assert float(pct.value) == 0.12
    assert pct.data_type == "n"
    assert "%" in pct.number_format

    pct_dec = by_label["pct_dec"]
    assert abs(float(pct_dec.value) - 0.125) < 1e-9
    assert "%" in pct_dec.number_format

    ident = by_label["id"]
    assert ident.value == "007"
    assert ident.data_type == "s"

    code = by_label["code"]
    assert code.value == "HARBOR-17"
    assert code.data_type == "s"

    formula = by_label["formula"]
    assert formula.value == "=1+2"
    assert formula.data_type == "s"


def test_xlsx_keeps_identifier_cells_as_text():
    from openpyxl import load_workbook

    body = """## Values

| label | value |
| --- | --- |
| zip | 32202 |
| phone | 9045551234 |
| e164 | +19045551234 |
| card | 4111111111111111 |
| account19 | 1234567890123456789 |
| isbn | 9780306406157 |
| mixed | $12% |
| grouped | 1,234 |
| money | $1,234.56 |
| pct | 12.5% |
| neg | -$5 |
| euro | €12.5 |
| pounds | £1,000 |
| lead0 | 007 |
| sci | 1e2 |
"""
    out = get_renderer("xlsx").render(title="Ids", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    ws = next(s for s in wb.worksheets if s.title != "Notes")
    by_label = {str(row[0].value): row[1] for row in ws.iter_rows(min_row=2) if row[0].value is not None}

    assert by_label["zip"].value == 32202
    assert by_label["zip"].data_type == "n"
    phone = by_label["phone"]
    assert phone.value == "9045551234"
    assert phone.data_type == "s"
    e164 = by_label["e164"]
    assert e164.value == "+19045551234"
    assert e164.data_type == "s"
    card = by_label["card"]
    assert card.value == "4111111111111111"
    assert card.data_type == "s"
    account19 = by_label["account19"]
    assert account19.value == "1234567890123456789"
    assert account19.data_type == "s"
    isbn = by_label["isbn"]
    assert isbn.value == "9780306406157"
    assert isbn.data_type == "s"
    mixed = by_label["mixed"]
    assert mixed.value == "$12%"
    assert mixed.data_type == "s"
    assert by_label["grouped"].value == 1234
    assert by_label["grouped"].data_type == "n"
    assert float(by_label["money"].value) == 1234.56
    assert abs(float(by_label["pct"].value) - 0.125) < 1e-9
    assert by_label["neg"].value == -5
    assert by_label["neg"].data_type == "n"
    assert "$" in by_label["neg"].number_format
    assert float(by_label["euro"].value) == 12.5
    assert "€" in by_label["euro"].number_format
    assert by_label["pounds"].value == 1000
    assert "£" in by_label["pounds"].number_format
    assert by_label["lead0"].value == "007"
    assert by_label["sci"].value == "1e2"


def test_xlsx_identifier_headers_keep_column_text():
    from openpyxl import load_workbook

    body = """## Contacts

| ZIP | Phone | Account | Amount | ISBN | Card |
| --- | --- | --- | --- | --- | --- |
| 32202 | 9045551234 | 1234567890123456789 | 1,234 | 9780306406157 | 4111111111111111 |
| 02134 | +19045551234 | 007 | $50.00 | 9780306406157 | 4111111111111111 |
"""
    out = get_renderer("xlsx").render(title="Contacts", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    ws = next(s for s in wb.worksheets if s.title != "Notes")
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers[:6] == ["ZIP", "Phone", "Account", "Amount", "ISBN", "Card"]
    row = next(ws.iter_rows(min_row=2, max_row=2))
    assert row[0].value == "32202" and row[0].data_type == "s"
    assert row[1].value == "9045551234" and row[1].data_type == "s"
    assert row[2].value == "1234567890123456789" and row[2].data_type == "s"
    assert row[3].value == 1234 and row[3].data_type == "n"
    assert row[4].value == "9780306406157" and row[4].data_type == "s"
    assert row[5].value == "4111111111111111" and row[5].data_type == "s"
    row2 = next(ws.iter_rows(min_row=3, max_row=3))
    assert row2[0].value == "02134" and row2[0].data_type == "s"
    assert row2[1].value == "+19045551234" and row2[1].data_type == "s"
    assert row2[3].value == 50 and row2[3].data_type == "n"


def test_xlsx_cells_parse_inline_markdown_only():
    from openpyxl import load_workbook

    body = """## Values

| label | value |
| --- | --- |
| spaced_money | - $5 |
| spaced_int | - 5 |
| ask | 1. Ask |
| quote | > kept |
| heading | # Title |
| bold_neg | - **$5** |
"""
    out = get_renderer("xlsx").render(title="Inline", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    ws = next(s for s in wb.worksheets if s.title != "Notes")
    by_label = {str(row[0].value): row[1] for row in ws.iter_rows(min_row=2) if row[0].value is not None}

    money = by_label["spaced_money"]
    assert money.value == -5
    assert money.data_type == "n"
    assert "$" in money.number_format

    spaced = by_label["spaced_int"]
    assert spaced.value == -5
    assert spaced.data_type == "n"

    ask = by_label["ask"]
    assert ask.value == "1. Ask"
    assert ask.data_type == "s"

    quote = by_label["quote"]
    assert quote.value == "> kept"
    assert quote.data_type == "s"

    heading = by_label["heading"]
    assert heading.value == "# Title"
    assert heading.data_type == "s"

    bold_neg = by_label["bold_neg"]
    assert bold_neg.value == -5
    assert bold_neg.data_type == "n"


def test_xlsx_count_headers_stay_numeric():
    from openpyxl import load_workbook

    body = """## Counts

| Number of invoices | No. of units | # of orders | Amount | Number | Account number | Card no. | Order # |
| --- | --- | --- | --- | --- | --- | --- | --- |
| 12 | 3 | 7 | 1,234 | 99 | 1234567890123456789 | 4111111111111111 | 42 |
"""
    out = get_renderer("xlsx").render(title="Counts", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    ws = next(s for s in wb.worksheets if s.title != "Notes")
    row = next(ws.iter_rows(min_row=2, max_row=2))
    assert row[0].value == 12 and row[0].data_type == "n"
    assert row[1].value == 3 and row[1].data_type == "n"
    assert row[2].value == 7 and row[2].data_type == "n"
    assert row[3].value == 1234 and row[3].data_type == "n"
    assert row[4].value == "99" and row[4].data_type == "s"
    assert row[5].value == "1234567890123456789" and row[5].data_type == "s"
    assert row[6].value == "4111111111111111" and row[6].data_type == "s"
    assert row[7].value == "42" and row[7].data_type == "s"


def test_xlsx_uses_model_sheet_names():
    from openpyxl import load_workbook

    long_name = "Quarterly revenue forecast extra words beyond limit"
    body = f"""## Q1 Revenue

| month | amount |
| --- | --- |
| Jan | 100 |

## Q1 Revenue

| month | amount |
| --- | --- |
| Feb | 200 |

## Foo/Bar:*Baz?

| a | b |
| --- | --- |
| 1 | 2 |

## {long_name}

| a | b |
| --- | --- |
| 3 | 4 |
"""
    out = get_renderer("xlsx").render(title="Book", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    names = [n for n in wb.sheetnames if n != "Notes"]
    assert names[0] == "Q1 Revenue"
    assert names[1] == "Q1 Revenue_2"
    assert names[2] == "FooBarBaz"
    assert names[3] == long_name[:31]
    assert all(len(n) <= 31 for n in wb.sheetnames)
    assert "month" not in wb.sheetnames
    assert "Foo/Bar:*Baz?" not in wb.sheetnames


def test_xlsx_sheet_names_follow_excel_rules():
    from openpyxl import load_workbook

    apostrophe_heading = ("A" * 30) + "'b"
    body = f"""## History

| a | b |
| --- | --- |
| 1 | 2 |

## history

| a | b |
| --- | --- |
| 3 | 4 |

## Notes

| a | b |
| --- | --- |
| 5 | 6 |

## {apostrophe_heading}

| a | b |
| --- | --- |
| 7 | 8 |

## Sales

| a | b |
| --- | --- |
| 9 | 10 |

## sales

| a | b |
| --- | --- |
| 11 | 12 |
"""
    out = get_renderer("xlsx").render(title="Book", body=body)
    wb = load_workbook(io.BytesIO(out.files["document.xlsx"]))
    names = wb.sheetnames
    assert "Notes" in names
    data_names = [n for n in names if n != "Notes"]
    assert "History" not in names
    assert data_names[0].lower() != "history"
    assert data_names[0].startswith("History")
    assert all(n.lower() != "history" for n in names)
    assert "Notes_2" in names
    assert names.count("Notes") == 1
    clipped = next(n for n in data_names if n.startswith("A"))
    assert len(clipped) <= 31
    assert not clipped.startswith("'")
    assert not clipped.endswith("'")
    assert all(len(n) <= 31 for n in names)
    assert all(not re.search(r"[\[\]*:/\\?]", n) for n in names)
    lowered = [n.lower() for n in names]
    assert len(lowered) == len(set(lowered))
    assert all(n.strip() for n in names)


def test_xlsx_builds_one_markdown_parser_per_file(monkeypatch: pytest.MonkeyPatch) -> None:
    from artifactsmith.renderers import md_parser

    calls = {"n": 0}
    orig = md_parser._new_linkify_parser

    def counting() -> object:
        calls["n"] += 1
        return orig()

    monkeypatch.setattr(md_parser, "_new_linkify_parser", counting)
    cols = 7
    rows = 34
    header = "| " + " | ".join(f"h{i}" for i in range(cols)) + " |"
    sep = "| " + " | ".join("---" for _ in range(cols)) + " |"
    data = "\n".join("| " + " | ".join(f"c{r}-{c}" for c in range(cols)) + " |" for r in range(rows))
    body = f"# Workbook\n\n## Inventory\n\n{header}\n{sep}\n{data}\n"
    out = get_renderer("xlsx").render(title="Workbook", body=body)
    assert calls["n"] == 1
    assert out.primary == "document.xlsx"
    assert out.files["document.xlsx"][:2] == b"PK"
