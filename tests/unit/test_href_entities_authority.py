"""Entity-like query keys must not be entity-decoded; authority confusion fails closed."""

from __future__ import annotations

import io
import re
import zipfile
import zlib
from html import escape

import pytest

from artifactsmith import builder as b
from artifactsmith.renderers import get_renderer
from artifactsmith.renderers.html_sanitize import sanitize_html_document
from artifactsmith.renderers.links import classify_href, emit_href, public_href_or_none
from artifactsmith.renderers.md_sanitize import collect_link_destinations, sanitize_markdown
from artifactsmith.renderers.safety import sanitize_text

# HTML5 named-character-reference prefixes that commonly appear as query keys.
# Unescaping would corrupt ``&{key}=`` into a Unicode character mid-query.
_ENTITY_QUERY_KEYS = [
    "section",
    "copy",
    "para",
    "currency",
    "timestamp",
    "region",
    "not",
    "or",
    "and",
    "pound",
    "euro",
    "yen",
    "cent",
    "times",
    "divide",
    "plusmn",
    "middot",
    "nbsp",
    "iexcl",
    "cent",
    "pound",
    "curren",
    "yen",
    "brvbar",
    "sect",  # &sect= → §
    "uml",
    "die",
    "acute",
    "micro",
    "para",
    "middot",
    "cedil",
    "sup1",
    "ordm",
    "raquo",
    "frac14",
    "frac12",
    "frac34",
    "iquest",
]

_AUTHORITY_CONFUSION = [
    "https://example.com%23@127.0.0.1/x",
    "https://example.com%3F@127.0.0.1/x",
    "https://example.com%2f@127.0.0.1/x",
    "https://example.com%2F@127.0.0.1/x",
    "https://user:pass@example.com/x",
    "https://user@example.com/x",
    r"https://example.com\@127.0.0.1/x",
    "https://[fe80::1%25eth0]/",
    "https://[::1%25eth0]/path",
]


def _pdf_uris(pdf: bytes) -> list[str]:
    uris: list[str] = []
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", pdf, re.S):
        try:
            dec = zlib.decompress(m.group(1))
        except zlib.error:
            continue
        for u in re.findall(rb"/URI\s*\((?:\\.|[^\\)])*\)", dec):
            inner = u[u.find(b"(") + 1 : -1]
            text = (
                inner.replace(b"\\(", b"(")
                .replace(b"\\)", b")")
                .replace(b"\\\\", b"\\")
                .decode("utf-8", errors="replace")
            )
            uris.append(text)
    return uris


@pytest.mark.parametrize("key", sorted(set(_ENTITY_QUERY_KEYS)))
def test_emit_preserves_entity_like_query_keys(key):
    url = f"https://example.com/search?a=1&{key}=value&z=9"
    assert "&" + key + "=" in url
    assert emit_href(url) == url
    assert public_href_or_none(url) == url
    # Must not contain the Unicode characters html.unescape would inject.
    assert "§" not in (public_href_or_none(url) or "")
    assert "©" not in (public_href_or_none(url) or "")
    assert "¶" not in (public_href_or_none(url) or "")
    assert "¤" not in (public_href_or_none(url) or "")
    assert "×" not in (public_href_or_none(url) or "")
    assert "®" not in (public_href_or_none(url) or "")


@pytest.mark.parametrize("key", sorted(set(_ENTITY_QUERY_KEYS)))
def test_entity_query_keys_survive_all_formats(key):
    url = f"https://example.com/search?a=1&{key}=value&z=9"
    needle = f"&{key}="
    md_out = sanitize_markdown(f"[cite]({url})")
    assert collect_link_destinations(md_out) == [url]
    assert needle in md_out

    href_attr = escape(url, quote=True)
    html = sanitize_html_document(
        "<!DOCTYPE html><html><head><meta charset='utf-8'><title>t</title></head>"
        f"<body><p><a href={href_attr}>cite</a></p></body></html>"
    )
    assert f'href="{escape(url, quote=True)}"' in html
    assert needle in html or escape(needle, quote=True) in html

    body = sanitize_text(f"See [cite]({url}) end.")
    pdf = get_renderer("pdf").render(title="T", body=body).files["document.pdf"]
    assert any(needle in u for u in _pdf_uris(pdf)), _pdf_uris(pdf)

    docx = get_renderer("docx").render(title="T", body=body).files["document.docx"]
    with zipfile.ZipFile(io.BytesIO(docx)) as zf:
        rels = "\n".join(zf.read(n).decode("utf-8", errors="replace") for n in zf.namelist() if n.endswith(".rels"))
    assert needle in rels or needle.replace("&", "&amp;") in rels

    from openpyxl import load_workbook

    xlsx = get_renderer("xlsx").render(title="T", body=body).files["document.xlsx"]
    wb = load_workbook(io.BytesIO(xlsx))
    targets: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink is not None:
                    targets.append(getattr(cell.hyperlink, "target", None) or str(cell.hyperlink))
    assert any(needle in t for t in targets), targets


@pytest.mark.parametrize("raw", _AUTHORITY_CONFUSION)
def test_authority_confusion_dropped(raw):
    assert classify_href(raw) == "blocked"
    assert public_href_or_none(raw) is None
    out = sanitize_markdown(f"[label]({raw})")
    assert collect_link_destinations(out) == []
    assert "label" in out


@pytest.mark.asyncio
async def test_builder_preserves_entity_query_and_drops_authority_confusion(monkeypatch):
    """Scripted model: keep &section= citations; neutralize authority-confusion URLs."""

    model_md = """# Brief

Cite [NSF search](https://www.nsf.gov/search?q=harbor&section=awards&region=us).

Ignore [trap](https://example.com%23@evil.example/secret).

Code is HARBOR-17.
"""

    async def fake_call(model, system, user, timeout=300):
        return (
            "===ASSUMPTIONS===\n- none\n===SUMMARY===\nBrief with citations.\n"
            f"===FILE: content.md===\n{model_md}\n===END===\n"
        )

    monkeypatch.setattr(b.llm, "call", fake_call)
    good = "https://www.nsf.gov/search?q=harbor&section=awards&region=us"
    for fmt in ("markdown", "pdf", "docx", "xlsx"):
        res = await b.run_build(
            kind="web_static",
            slug="brief",
            display_name="Brief",
            verbatim="Cite NSF search with section and region; state HARBOR-17.",
            model="test-model",
            base_source=None,
            base_version=None,
            history=[],
            progress=lambda _m: None,
            source={
                "source_content": (
                    "Pilot store code is HARBOR-17. NSF search: "
                    "https://www.nsf.gov/search?q=harbor&section=awards&region=us"
                )
            },
            format=fmt,
        )
        if fmt == "markdown":
            text = res.files["document.md"].decode()
            assert good in text
            assert "&section=" in text
            assert "&region=" in text
            assert "evil.example" not in collect_link_destinations(text)
            assert collect_link_destinations(text) == [good]
        else:
            md = res.files["content.md"].decode()
            assert good in md
            assert collect_link_destinations(md) == [good]
            if fmt == "pdf":
                uris = _pdf_uris(res.files["document.pdf"])
                assert good in uris
                assert not any("evil.example" in u for u in uris)
            elif fmt == "docx":
                with zipfile.ZipFile(io.BytesIO(res.files["document.docx"])) as zf:
                    rels = "\n".join(
                        zf.read(n).decode("utf-8", errors="replace") for n in zf.namelist() if n.endswith(".rels")
                    )
                assert "section=awards" in rels
                assert "evil.example" not in rels
            else:
                from openpyxl import load_workbook

                wb = load_workbook(io.BytesIO(res.files["document.xlsx"]))
                targets: list[str] = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows():
                        for cell in row:
                            if cell.hyperlink is not None:
                                targets.append(getattr(cell.hyperlink, "target", None) or str(cell.hyperlink))
                assert any("&section=" in t for t in targets), targets
                assert not any("evil.example" in t for t in targets)
