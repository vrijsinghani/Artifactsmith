"""DOCX renderer via python-docx (MIT)."""

from __future__ import annotations

import io

from .base import RenderOutput
from .md_parse import parse_blocks


class DocxRenderer:
    format = "docx"
    model_filename = "content.md"

    def render(self, *, title: str, body: str) -> RenderOutput:
        from docx import Document
        from docx.shared import Pt

        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        if title:
            doc.add_heading(title, level=0)
        for b in parse_blocks(body):
            if b.kind == "heading":
                doc.add_heading(b.text, level=min(max(b.level, 1), 4))
            elif b.kind == "paragraph":
                doc.add_paragraph(b.text)
            elif b.kind == "list":
                for item in b.items:
                    doc.add_paragraph(item, style="List Bullet")
            elif b.kind == "code":
                p = doc.add_paragraph(b.text)
                for run in p.runs:
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)
            elif b.kind == "table":
                cols = max(len(b.headers), max((len(r) for r in b.rows), default=0))
                if cols == 0:
                    continue
                table = doc.add_table(rows=1 + len(b.rows), cols=cols)
                table.style = "Table Grid"
                hdr = table.rows[0].cells
                for i, h in enumerate(b.headers):
                    if i < cols:
                        hdr[i].text = h
                for ri, row in enumerate(b.rows):
                    cells = table.rows[ri + 1].cells
                    for ci, val in enumerate(row):
                        if ci < cols:
                            cells[ci].text = val
        buf = io.BytesIO()
        doc.save(buf)
        files = {
            "document.docx": buf.getvalue(),
            "content.md": (body.strip() + "\n").encode("utf-8"),
        }
        return RenderOutput(files=files, primary="document.docx")
