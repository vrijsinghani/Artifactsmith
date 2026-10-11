"""XLSX renderer via openpyxl (MIT). Tables become sheets; prose becomes a Notes sheet.

Numeric values are stored as numbers (currency/percent formats). Other cells are
string literals (never formulas). Public http(s) URLs get a plain hyperlink
relationship, not ``=HYPERLINK()``. Sheet names come from the heading above each table.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence
from typing import Any

from .base import RenderOutput
from .links import iter_inline_segments, public_href_or_none
from .md_parse import parse_blocks, visible_text
from .safety import strip_ooxml_controls

_SHEET_SAFE = re.compile(r"[\[\]\*\:\/\\\?]")
_GROUPED_NUM = re.compile(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?")
_PLAIN_NUM = re.compile(r"\d+(?:\.\d+)?")
_CURRENCY_SYM = frozenset("$£€¥")
_EXCEL_NAME_MAX = 31


def _sheet_name(title: str, used: set[str], index: int) -> str:
    cleaned = strip_ooxml_controls(visible_text(title))
    base = _SHEET_SAFE.sub("", cleaned).strip().strip("'") or f"Sheet{index}"
    name = base[:_EXCEL_NAME_MAX]
    n = 2
    while name.lower() in used or not name:
        suffix = f"_{n}"
        name = (base[: max(0, _EXCEL_NAME_MAX - len(suffix))] + suffix)[:_EXCEL_NAME_MAX]
        n += 1
    used.add(name.lower())
    return name


def _currency_format(symbol: str, decimals: int) -> str:
    dec = f".{'0' * decimals}" if decimals else ""
    if symbol == "$":
        return f"$#,##0{dec}"
    return f'"{symbol}"#,##0{dec}'


def _parse_numeric(raw: str) -> tuple[int | float, str] | None:
    """Return (number, Excel format) for ints, decimals, currency, and percents.

    Commas are thousands separators. Leading-zero IDs and formula-like values stay text.
    """
    text = raw.strip()
    if not text or text[0] in "=@":
        return None
    percent = text.endswith("%")
    if percent:
        text = text[:-1].strip()
    sign = 1
    currency = ""
    if text and text[0] in "+-":
        sign = -1 if text[0] == "-" else 1
        text = text[1:].strip()
    if text and text[0] in _CURRENCY_SYM:
        currency = text[0]
        text = text[1:].strip()
    if text and text[0] in "+-":
        sign = -1 if text[0] == "-" else 1
        text = text[1:].strip()
    if not text:
        return None
    if _GROUPED_NUM.fullmatch(text):
        grouped = True
        core = text.replace(",", "")
    elif _PLAIN_NUM.fullmatch(text):
        grouped = False
        core = text
    else:
        return None
    if re.match(r"^0\d", core):
        return None
    if "." in core:
        number: int | float = sign * float(core)
        decimals = len(core.split(".", 1)[1])
    else:
        number = sign * int(core)
        decimals = 0
    if percent:
        return float(number) / 100.0, ("0%" if decimals == 0 else f"0.{'0' * decimals}%")
    if currency:
        return number, _currency_format(currency, decimals)
    if decimals:
        return number, (f"#,##0.{'0' * decimals}" if grouped else f"0.{'0' * decimals}")
    return number, ("#,##0" if grouped else "0")


def _first_public_href(text: str) -> str | None:
    for _display, href in iter_inline_segments(text):
        if href:
            return href
    return public_href_or_none(text.strip())


def _literal_cell(ws: Any, row: int, col: int, value: object) -> None:
    """Write a cell: numbers as numbers, everything else as a string (never a formula).

    The full cell text is preserved (every URL stays visible). When multiple public
    http(s) URLs appear, only the first becomes the cell hyperlink.
    """
    from openpyxl.cell.cell import TYPE_STRING

    text = strip_ooxml_controls("" if value is None else str(value))
    parsed = _parse_numeric(text)
    cell = ws.cell(row=row, column=col)
    if parsed is not None:
        number, fmt = parsed
        cell.value = number
        cell.number_format = fmt
        return
    cell.value = text
    cell.data_type = TYPE_STRING
    href = _first_public_href(text) if text else None
    if href:
        # Plain hyperlink relationship — never a formula.
        cell.hyperlink = href


def _write_row(ws: Any, row: int, values: Sequence[object]) -> None:
    for col, value in enumerate(values, start=1):
        _literal_cell(ws, row, col, value)


class XlsxRenderer:
    format = "xlsx"
    model_filename = "content.md"

    def render(self, *, title: str, body: str) -> RenderOutput:
        from openpyxl import Workbook

        wb = Workbook()
        # Remove the default sheet; we recreate intentionally.
        default = wb.active
        if default is not None:
            wb.remove(default)
        used: set[str] = set()
        notes: list[str] = []
        table_i = 0
        pending_sheet: str | None = None
        title = strip_ooxml_controls(title)
        if title:
            notes.append(title)
        for b in parse_blocks(body):
            if b.kind == "heading":
                pending_sheet = b.text
                notes.append(b.text)
                continue
            if b.kind == "table":
                table_i += 1
                raw_name = pending_sheet or (b.headers[0] if b.headers else f"Table{table_i}")
                pending_sheet = None
                name = _sheet_name(raw_name, used, table_i)
                ws = wb.create_sheet(name)
                headers = b.headers or [f"col{i + 1}" for i in range(max((len(r) for r in b.rows), default=1))]
                _write_row(ws, 1, list(headers))
                width = len(headers)
                for i, row in enumerate(b.rows, start=2):
                    padded = list(row) + [""] * max(0, width - len(row))
                    _write_row(ws, i, padded[:width])
            elif b.kind == "paragraph":
                notes.append(b.text)
            elif b.kind == "list":
                notes.extend(f"- {it}" for it in b.items)
            elif b.kind == "code":
                notes.append(b.text)
        if not wb.sheetnames:
            ws = wb.create_sheet(_sheet_name(title or "Data", used, 1))
            _write_row(ws, 1, ["content"])
            _write_row(ws, 2, [body.strip() or "(empty)"])
        if notes:
            nws = wb.create_sheet(_sheet_name("Notes", used, 99))
            _write_row(nws, 1, ["notes"])
            for i, line in enumerate(notes, start=2):
                _write_row(nws, i, [line])
        buf = io.BytesIO()
        wb.save(buf)
        files = {
            "document.xlsx": buf.getvalue(),
            "content.md": (body.strip() + "\n").encode("utf-8"),
        }
        return RenderOutput(files=files, primary="document.xlsx")
